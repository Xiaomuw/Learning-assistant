from __future__ import annotations

import csv
import json
import re
import uuid
from pathlib import Path
from typing import Any

from app.infrastructure.db.repositories.practice import upsert_question_bank_item_records
from app.knowledge.concept_registry import get_concept_registry
from app.services.rag.indexing import index_question_bank_items


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return [dict(row) for row in reader]


def _read_json_rows(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [dict(item) for item in payload]
    if isinstance(payload, dict) and isinstance(payload.get("items"), list):
        return [dict(item) for item in payload["items"]]
    raise ValueError("JSON 题库文件格式不正确，应为数组或 {\"items\": [...]}。")


def _read_excel_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("导入 Excel 题库需要安装 openpyxl。") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    data_rows: list[dict[str, Any]] = []
    for row in rows[1:]:
        data_rows.append({
            headers[index]: row[index]
            for index in range(min(len(headers), len(row)))
            if headers[index]
        })
    return data_rows


def _coerce_import_row(raw: dict[str, Any], default_subject_id: str | None = None) -> dict[str, Any]:
    options = raw.get("options")
    if isinstance(options, str):
        try:
            options = json.loads(options)
        except json.JSONDecodeError:
            options = [
                {"key": "A", "text": str(raw.get("option_a", "") or "").strip()},
                {"key": "B", "text": str(raw.get("option_b", "") or "").strip()},
                {"key": "C", "text": str(raw.get("option_c", "") or "").strip()},
                {"key": "D", "text": str(raw.get("option_d", "") or "").strip()},
            ]

    rationales = raw.get("option_rationales")
    if isinstance(rationales, str):
        try:
            rationales = json.loads(rationales)
        except json.JSONDecodeError:
            rationales = {
                "A": str(raw.get("rationale_a", "") or "").strip(),
                "B": str(raw.get("rationale_b", "") or "").strip(),
                "C": str(raw.get("rationale_c", "") or "").strip(),
                "D": str(raw.get("rationale_d", "") or "").strip(),
            }

    tags = raw.get("tags", [])
    if isinstance(tags, str):
        tags = [tag.strip() for tag in re.split(r"[,，;；]", tags) if tag.strip()]

    item = {
        "id": str(raw.get("id") or raw.get("question_id") or uuid.uuid4()),
        "subject_id": str(raw.get("subject_id") or default_subject_id or "general"),
        "knowledge_point": str(raw.get("knowledge_point") or "").strip(),
        "question": str(raw.get("question") or raw.get("stem") or "").strip(),
        "options": options,
        "correct_option": str(raw.get("correct_option") or raw.get("answer") or "A").strip().upper()[:1],
        "explanation": str(raw.get("explanation") or "").strip(),
        "option_rationales": rationales,
        "source_name": str(raw.get("source_name") or raw.get("exam_name") or "").strip(),
        "source_year": int(raw["source_year"]) if str(raw.get("source_year", "")).strip() else None,
        "difficulty": str(raw.get("difficulty") or "基础").strip() or "基础",
        "tags": tags,
    }
    item.update(_concept_fields_for_import_item(item, raw))
    return item


def _concept_fields_for_import_item(item: dict[str, Any], raw: dict[str, Any]) -> dict[str, str]:
    explicit_concept_id = str(raw.get("concept_id") or raw.get("primary_concept_id") or "").strip()
    registry = get_concept_registry()
    if explicit_concept_id:
        concept = registry.get_concept(explicit_concept_id)
        if concept is not None:
            return {
                "chapter_id": concept.chapter_id,
                "topic_id": concept.topic_id,
                "concept_id": concept.concept_id,
                "concept_name": concept.concept_name,
                "concept_mapping_status": "matched",
            }
    match = registry.match_query(
        "\n".join(
            str(item.get(key) or "")
            for key in ("knowledge_point", "question", "explanation")
            if str(item.get(key) or "").strip()
        ),
        subject_id=str(item.get("subject_id") or ""),
    )
    if match.status == "matched" and match.concept_id:
        concept = registry.get_concept(match.concept_id)
        if concept is not None:
            return {
                "chapter_id": concept.chapter_id,
                "topic_id": concept.topic_id,
                "concept_id": concept.concept_id,
                "concept_name": concept.concept_name,
                "concept_mapping_status": "matched",
            }
    return {
        "chapter_id": str(raw.get("chapter_id") or ""),
        "topic_id": str(raw.get("topic_id") or ""),
        "concept_id": explicit_concept_id,
        "concept_name": str(raw.get("concept_name") or ""),
        "concept_mapping_status": match.status,
    }


def import_question_bank_file(path: str, *, default_subject_id: str | None = None) -> int:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv_rows(file_path)
    elif suffix == ".json":
        rows = _read_json_rows(file_path)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _read_excel_rows(file_path)
    else:
        raise ValueError("暂不支持该题库文件类型，请使用 CSV / JSON / Excel。")

    items = [_coerce_import_row(row, default_subject_id=default_subject_id) for row in rows]
    saved = upsert_question_bank_item_records(items)
    indexed_items = [
        {
            **item,
            "question_id": item["id"],
            "source_type": "question_bank",
        }
        for item in items
    ]
    index_question_bank_items(indexed_items)
    return saved
